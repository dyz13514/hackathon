"""Strict, data-driven import of a complete planning workbook.

The workbook supplies every scheduling input. This module contains a schema,
not demo entity values; nothing is seeded or inferred from a product ID.
"""

from __future__ import annotations

import io
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import audit
from app.db import models as orm
from app.services.ingestion import ImportDataError
from app.services.spreadsheet import MAX_FILE_BYTES, MAX_ROWS


@dataclass(frozen=True)
class SheetSpec:
    model: type
    columns: tuple[str, ...]
    required: tuple[str, ...]
    key: tuple[str, ...]


# Dependency order matters when committing to a database with foreign keys.
SHEETS: dict[str, SheetSpec] = {
    "Products": SheetSpec(
        orm.Product, ("product_id", "name", "description"),
        ("product_id", "name"), ("product_id",),
    ),
    "Materials": SheetSpec(
        orm.Material,
        ("material_id", "name", "unit", "quantity_available", "reserved_quantity"),
        ("material_id", "name", "unit", "quantity_available", "reserved_quantity"),
        ("material_id",),
    ),
    "Machines": SheetSpec(
        orm.Machine,
        ("machine_id", "machine_type", "capabilities", "status", "available_start",
         "available_end", "rate_multiplier"),
        ("machine_id", "machine_type", "capabilities", "status", "available_start",
         "available_end", "rate_multiplier"),
        ("machine_id",),
    ),
    "Workers": SheetSpec(
        orm.Worker, ("worker_id", "name", "skills", "shift_start", "shift_end"),
        ("worker_id", "name", "skills", "shift_start", "shift_end"), ("worker_id",),
    ),
    "Operations": SheetSpec(
        orm.Operation,
        ("operation_id", "product_id", "sequence", "required_machine_type",
         "required_capability", "required_worker_skill", "base_processing_time_per_unit",
         "setup_time"),
        ("operation_id", "product_id", "sequence", "required_machine_type",
         "required_worker_skill", "base_processing_time_per_unit", "setup_time"),
        ("operation_id",),
    ),
    "ProductMaterials": SheetSpec(
        orm.ProductMaterial, ("product_id", "material_id", "quantity_per_unit"),
        ("product_id", "material_id", "quantity_per_unit"),
        ("product_id", "material_id"),
    ),
    "IncomingDeliveries": SheetSpec(
        orm.IncomingDelivery, ("delivery_id", "material_id", "quantity", "eta", "confirmed"),
        ("delivery_id", "material_id", "quantity", "eta", "confirmed"),
        ("delivery_id",),
    ),
    "MachineDowntime": SheetSpec(
        orm.MachineDowntime, ("downtime_id", "machine_id", "start_time", "end_time", "reason"),
        ("downtime_id", "machine_id", "start_time", "end_time", "reason"),
        ("downtime_id",),
    ),
    "WorkerAbsences": SheetSpec(
        orm.WorkerAbsence, ("absence_id", "worker_id", "start_time", "end_time"),
        ("absence_id", "worker_id", "start_time", "end_time"),
        ("absence_id",),
    ),
    "ChangeoverRules": SheetSpec(
        orm.ChangeoverRule,
        ("rule_id", "machine_id", "from_product_id", "to_product_id",
         "changeover_minutes", "specificity"),
        ("rule_id", "changeover_minutes", "specificity"), ("rule_id",),
    ),
    "Orders": SheetSpec(
        orm.Order,
        ("order_id", "product_id", "quantity", "due_date", "promised_date",
         "priority", "notes"),
        ("order_id", "product_id", "quantity", "due_date", "priority"),
        ("order_id",),
    ),
}

_DATES = {
    "available_start", "available_end", "shift_start", "shift_end", "eta",
    "start_time", "end_time", "due_date", "promised_date",
}
_DECIMALS = {
    "quantity_available", "reserved_quantity", "rate_multiplier",
    "base_processing_time_per_unit", "quantity_per_unit", "quantity",
}
_INTEGERS = {"sequence", "setup_time", "changeover_minutes", "specificity"}
_LISTS = {"capabilities", "skills"}
_SOURCES = {"Products", "Materials", "Machines", "Workers", "Orders", "IncomingDeliveries"}
_ROOTS = {"Products", "Materials", "Machines", "Workers", "Orders"}


@dataclass(frozen=True)
class PackageRow:
    row_number: int
    values: dict[str, Any]
    raw_row: list[str]


@dataclass(frozen=True)
class PlanningPackage:
    sheets: dict[str, list[PackageRow]]
    anchor_date: date | None

    @property
    def counts(self) -> dict[str, int]:
        return {name: len(rows) for name, rows in self.sheets.items()}

    @property
    def row_count(self) -> int:
        return sum(self.counts.values())


def parse_planning_package(*, filename: str, content: bytes) -> PlanningPackage:
    """Parse and fully validate a canonical workbook without touching the database."""
    if not filename.lower().endswith(".xlsx") or content[:2] != b"PK":
        raise ImportDataError("A complete planning package must be a valid .xlsx workbook.")
    if len(content) > MAX_FILE_BYTES:
        raise ImportDataError("Planning workbook exceeds the 5 MB upload limit.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if any("vbaProject.bin" in name for name in archive.namelist()):
                raise ImportDataError("Macro-enabled workbooks are not accepted.")
    except zipfile.BadZipFile as error:
        raise ImportDataError("The workbook is not a valid XLSX file.") from error

    import openpyxl

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise ImportDataError("The workbook could not be read.") from error
    try:
        missing = [name for name in SHEETS if name not in workbook.sheetnames]
        if missing:
            raise ImportDataError("Missing workbook sheets: " + ", ".join(missing))
        unexpected = set(workbook.sheetnames) - set(SHEETS) - {"Guide"}
        if unexpected:
            raise ImportDataError("Unexpected workbook sheets: " + ", ".join(sorted(unexpected)))
        sheets: dict[str, list[PackageRow]] = {}
        for name, spec in SHEETS.items():
            iterator = workbook[name].iter_rows(values_only=True)
            header = next(iterator, None)
            if header is None:
                raise ImportDataError(f"{name}: header row is missing.")
            names = [str(item).strip() if item is not None else "" for item in header]
            if names != list(spec.columns):
                raise ImportDataError(
                    f"{name}: columns must be exactly {', '.join(spec.columns)}."
                )
            entries: list[PackageRow] = []
            keys: set[tuple[str, ...]] = set()
            for row_no, raw in enumerate(iterator, start=2):
                if not any(item is not None and str(item).strip() for item in raw):
                    continue
                if len(raw) > len(spec.columns) and any(
                    item is not None and str(item).strip() for item in raw[len(spec.columns):]
                ):
                    raise ImportDataError(f"{name} row {row_no}: extra populated columns.")
                values = {
                    field: _parse_cell(
                        field, raw[index] if index < len(raw) else None, name, row_no
                    )
                    for index, field in enumerate(spec.columns)
                }
                for field in spec.required:
                    if values[field] is None or values[field] == "" or values[field] == []:
                        raise ImportDataError(f"{name} row {row_no}: {field} is required.")
                key = tuple(str(values[field]) for field in spec.key)
                if key in keys:
                    raise ImportDataError(f"{name} row {row_no}: duplicate key {key}.")
                keys.add(key)
                entries.append(
                    PackageRow(
                        row_no, values,
                        [_raw_text(item) for item in raw[:len(spec.columns)]],
                    )
                )
                if sum(len(part) for part in sheets.values()) + len(entries) > MAX_ROWS:
                    raise ImportDataError("Planning workbook exceeds the 2,000-row limit.")
            sheets[name] = entries
        _validate_relationships(sheets)
        order_dates = [row.values["due_date"].date() for row in sheets["Orders"]]
        return PlanningPackage(sheets, min(order_dates) if order_dates else None)
    finally:
        workbook.close()


def _raw_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    return str(value)[:500]


def _parse_cell(field: str, value: Any, sheet: str, row_no: int) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if field in _DATES:
        if isinstance(value, datetime):
            return value.replace(tzinfo=None)
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        try:
            return datetime.fromisoformat(str(value).strip())
        except ValueError as error:
            raise ImportDataError(f"{sheet} row {row_no}: invalid {field} datetime.") from error
    if field in _DECIMALS:
        try:
            number = Decimal(str(value).strip())
        except InvalidOperation as error:
            raise ImportDataError(f"{sheet} row {row_no}: invalid {field} number.") from error
        if not number.is_finite():
            raise ImportDataError(f"{sheet} row {row_no}: {field} must be finite.")
        return number
    if field in _INTEGERS:
        try:
            return int(str(value).strip())
        except ValueError as error:
            raise ImportDataError(f"{sheet} row {row_no}: invalid {field} integer.") from error
    if field in _LISTS:
        if not isinstance(value, str):
            raise ImportDataError(
                f"{sheet} row {row_no}: {field} must use semicolon-separated text."
            )
        return [item.strip() for item in value.split(";") if item.strip()]
    if field == "confirmed":
        normalized = str(value).strip().lower()
        if normalized in {"true", "yes", "1", "是"}:
            return True
        if normalized in {"false", "no", "0", "否"}:
            return False
        raise ImportDataError(f"{sheet} row {row_no}: confirmed must be true or false.")
    text = str(value).strip()
    if len(text) > 500:
        raise ImportDataError(f"{sheet} row {row_no}: {field} exceeds 500 characters.")
    return text


def _validate_relationships(sheets: dict[str, list[PackageRow]]) -> None:
    ids = {
        name: {str(row.values[SHEETS[name].key[0]]) for row in sheets[name]}
        for name in _ROOTS
    }
    for name in ("Products", "Machines", "Workers", "Orders"):
        if not sheets[name]:
            raise ImportDataError(f"{name}: at least one row is required for scheduling.")
    relations = {
        "Operations": (("product_id", "Products"),),
        "ProductMaterials": (("product_id", "Products"), ("material_id", "Materials")),
        "IncomingDeliveries": (("material_id", "Materials"),),
        "MachineDowntime": (("machine_id", "Machines"),),
        "WorkerAbsences": (("worker_id", "Workers"),),
        "ChangeoverRules": (("machine_id", "Machines"), ("from_product_id", "Products"),
                            ("to_product_id", "Products")),
        "Orders": (("product_id", "Products"),),
    }
    for sheet_name, links in relations.items():
        for row in sheets[sheet_name]:
            for field, target in links:
                value = row.values[field]
                if value is not None and str(value) not in ids[target]:
                    raise ImportDataError(
                        f"{sheet_name} row {row.row_number}: {field} {value} is not in {target}."
                    )
    routes: dict[str, set[int]] = {product_id: set() for product_id in ids["Products"]}
    for row in sheets["Operations"]:
        value = row.values
        sequence = value["sequence"]
        if sequence not in {1, 2, 3} or sequence in routes[value["product_id"]]:
            raise ImportDataError(
                f"Operations row {row.row_number}: sequence must be unique and between 1 and 3."
            )
        routes[value["product_id"]].add(sequence)
        if value["base_processing_time_per_unit"] <= 0 or value["setup_time"] < 0:
            raise ImportDataError(
                f"Operations row {row.row_number}: processing/setup time is invalid."
            )
    for product_id, sequence_set in routes.items():
        if sequence_set != set(range(1, len(sequence_set) + 1)):
            raise ImportDataError(
                f"Products {product_id}: routing must contain contiguous steps 1–3."
            )
    ordered_products = {row.values["product_id"] for row in sheets["Orders"]}
    for product_id in ordered_products:
        if not routes[product_id]:
            raise ImportDataError(
                f"Products {product_id}: ordered products need at least one operation."
            )
    machines = [row.values for row in sheets["Machines"]]
    worker_skills = {
        skill for row in sheets["Workers"] for skill in row.values["skills"]
    }
    for row in sheets["Operations"]:
        operation = row.values
        if not any(
            machine["machine_type"] == operation["required_machine_type"]
            and (
                operation["required_capability"] is None
                or operation["required_capability"] in machine["capabilities"]
            )
            for machine in machines
        ):
            raise ImportDataError(
                f"Operations row {row.row_number}: no machine has the required type/capability."
            )
        if operation["required_worker_skill"] not in worker_skills:
            raise ImportDataError(
                f"Operations row {row.row_number}: no worker has the required skill."
            )
    for sheet_name, start, end in (
        ("Machines", "available_start", "available_end"),
        ("Workers", "shift_start", "shift_end"),
        ("MachineDowntime", "start_time", "end_time"),
        ("WorkerAbsences", "start_time", "end_time"),
    ):
        for row in sheets[sheet_name]:
            if row.values[start] >= row.values[end]:
                raise ImportDataError(
                    f"{sheet_name} row {row.row_number}: {end} must follow {start}."
                )
    for row in sheets["Machines"]:
        if row.values["status"] not in orm.MACHINE_STATUSES:
            raise ImportDataError(f"Machines row {row.row_number}: invalid status.")
        if row.values["rate_multiplier"] <= 0:
            raise ImportDataError(
                f"Machines row {row.row_number}: rate_multiplier must be positive."
            )
    for row in sheets["Workers"]:
        if not row.values["skills"]:
            raise ImportDataError(f"Workers row {row.row_number}: skills cannot be empty.")
    for sheet_name, fields in (
        ("Materials", ("quantity_available", "reserved_quantity")),
        ("ProductMaterials", ("quantity_per_unit",)),
        ("IncomingDeliveries", ("quantity",)),
        ("Orders", ("quantity",)),
    ):
        for row in sheets[sheet_name]:
            for field in fields:
                minimum_invalid = (
                    row.values[field] < 0 if sheet_name == "Materials"
                    else row.values[field] <= 0
                )
                if minimum_invalid:
                    raise ImportDataError(f"{sheet_name} row {row.row_number}: {field} is invalid.")
    for row in sheets["Orders"]:
        if row.values["priority"] not in orm.ORDER_PRIORITIES:
            raise ImportDataError(f"Orders row {row.row_number}: invalid priority.")
    for row in sheets["MachineDowntime"]:
        if row.values["reason"] not in orm.DOWNTIME_REASONS:
            raise ImportDataError(f"MachineDowntime row {row.row_number}: invalid reason.")
    for row in sheets["ChangeoverRules"]:
        if row.values["changeover_minutes"] < 0 or row.values["specificity"] not in {1, 2, 3}:
            raise ImportDataError(f"ChangeoverRules row {row.row_number}: invalid rule.")


def commit_planning_package(
    session: Session, *, package: PlanningPackage, filename: str, checksum: str,
    now: datetime | None = None,
) -> str:
    """Commit a workbook into a fresh workspace or replay a reverted package."""
    existing_roots = {
        name: list(session.scalars(select(SHEETS[name].model)).all()) for name in _ROOTS
    }
    replay = any(existing_roots.values())
    if replay:
        previous = session.scalar(
            select(func.count()).select_from(orm.ImportBatch).where(
                orm.ImportBatch.entity_type == "PACKAGE",
                orm.ImportBatch.status == "REVERTED",
            )
        )
        if not previous:
            raise ImportDataError(
                "The planning workspace already contains data. Import a complete workbook "
                "into an empty workspace, or revert the previous package first."
            )
        for name, rows in existing_roots.items():
            expected = {
                tuple(str(item.values[field]) for field in SHEETS[name].key)
                for item in package.sheets[name]
            }
            actual = {
                tuple(str(getattr(row, field)) for field in SHEETS[name].key)
                for row in rows
            }
            if actual != expected or any(row.record_status != "REVERTED" for row in rows):
                raise ImportDataError(
                    f"{name}: replay requires the same IDs as the reverted package "
                    "and no active records."
                )
        for name in set(SHEETS) - _ROOTS - {"ChangeoverRules"}:
            spec = SHEETS[name]
            expected = {
                tuple(str(item.values[field]) for field in spec.key)
                for item in package.sheets[name]
            }
            actual = {
                tuple(str(getattr(row, field)) for field in spec.key)
                for row in session.scalars(select(spec.model)).all()
            }
            if actual != expected:
                raise ImportDataError(
                    f"{name}: replay requires the same row IDs as the reverted package."
                )
    resolved_now = now or datetime.now()
    batch_id = f"BATCH-{uuid.uuid4().hex[:12]}"
    schema = {name: list(spec.columns) for name, spec in SHEETS.items()}
    session.add(
        orm.ImportBatch(
            batch_id=batch_id, file_name=filename, file_checksum=checksum,
            entity_type="PACKAGE", row_count=package.row_count, status="COMMITTED",
            proposed_mapping={"kind": "EXACT_WORKBOOK_SCHEMA", "sheets": schema},
            accepted_mapping={"kind": "EXACT_WORKBOOK_SCHEMA", "sheets": schema},
            operator_decisions={"confirmed": True}, normalisations={}, unparsed_cells=[],
            formula_columns=[], ingestion_report={"sheets": package.counts},
            imported_at=resolved_now, created_at=resolved_now,
        )
    )
    session.flush()
    for name, spec in SHEETS.items():
        for item in package.sheets[name]:
            values = dict(item.values)
            if name in _SOURCES:
                values["source"] = "SPREADSHEET_IMPORT"
                values["last_updated_at"] = resolved_now
            if name in _ROOTS:
                values["record_status"] = "ACTIVE"
                values["import_batch_id"] = batch_id
            if name == "Orders":
                values["injection_suspected"] = False
                values["source_row_number"] = item.row_number
            if name in {"MachineDowntime", "WorkerAbsences"}:
                values["disruption_id"] = None
            entity = spec.model(**values)
            session.merge(entity)
            session.add(
                orm.ImportRowProvenance(
                    id=f"PROV-{uuid.uuid4().hex[:12]}", batch_id=batch_id,
                    entity_type=name.upper(),
                    entity_id="|".join(str(item.values[field]) for field in spec.key),
                    source_row_number=item.row_number, raw_row=item.raw_row,
                    overwritten_payload=None,
                )
            )
        session.flush()
    session.commit()
    audit.append(
        event_category="DATA_IMPORT", event_type="COMMIT_BATCH", actor="PLANNER",
        payload={"batch_id": batch_id, "entity_type": "PACKAGE", "sheet_counts": package.counts},
        subject_type="ImportBatch", subject_id=batch_id, occurred_at=resolved_now,
    )
    return batch_id


def revert_planning_package(
    session: Session, *, batch_id: str, now: datetime | None = None
) -> int:
    """Exclude every package root from future snapshots while keeping plan history."""
    batch = session.get(orm.ImportBatch, batch_id)
    if batch is None or batch.entity_type != "PACKAGE":
        raise ImportDataError("Planning workbook batch not found.")
    if batch.status == "REVERTED":
        return 0
    resolved_now = now or datetime.now()
    rows = list(
        session.scalars(
            select(orm.ImportRowProvenance).where(orm.ImportRowProvenance.batch_id == batch_id)
        ).all()
    )
    for item in rows:
        name = next((sheet for sheet in SHEETS if sheet.upper() == item.entity_type), None)
        if name is None:
            continue
        spec = SHEETS[name]
        key_values = item.entity_id.split("|")
        key: Any = key_values[0] if len(key_values) == 1 else tuple(key_values)
        entity = session.get(spec.model, key)
        if entity is None:
            continue
        if name in _ROOTS:
            if entity.import_batch_id == batch_id:
                entity.record_status = "REVERTED"
                entity.last_updated_at = resolved_now
        elif name == "ChangeoverRules":
            session.delete(entity)
    batch.status = "REVERTED"
    session.commit()
    audit.append(
        event_category="DATA_IMPORT", event_type="REVERT_BATCH", actor="PLANNER",
        payload={"batch_id": batch_id, "reverted_row_count": len(rows)},
        subject_type="ImportBatch", subject_id=batch_id, occurred_at=resolved_now,
    )
    return len(rows)
