from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "AI生产排产副驾商业计划书.docx"

BLUE = "173F5F"
LIGHT_BLUE = "EAF2F8"
PALE = "F6F8FA"
GRAY = "D9D9D9"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def borders(table):
    tbl_pr = table._tbl.tblPr
    el = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = OxmlElement(f'w:{edge}')
        tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'6'); tag.set(qn('w:color'),GRAY)
        el.append(tag)
    tbl_pr.append(el)

def cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    borders(table)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; shade(cell, BLUE); cell_text(cell, h, True, 'FFFFFF')
        if widths: cell.width = Inches(widths[i])
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if ri % 2: shade(cells[i], PALE)
            cell_text(cells[i], str(val), size=9)
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def para(doc, text='', bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def bullets(doc, items):
    for x in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3); p.add_run(x)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.72); sec.bottom_margin = Inches(.72)
sec.left_margin = Inches(.78); sec.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
styles['Normal'].font.size = Pt(10.5)
for name, size, color in [('Title',30,'000000'),('Heading 1',17,'000000'),('Heading 2',12,'000000')]:
    s=styles[name]; s.font.name='Aptos Display'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
    s.font.bold=True
    s.paragraph_format.space_before=Pt(15); s.paragraph_format.space_after=Pt(8)

# Cover
p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('AI生产排产副驾商业计划书')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(14)
r=p.add_run('面向中小制造车间的可解释、可审批、可审计生产决策平台'); r.font.size=Pt(16); r.font.color.rgb=RGBColor.from_string(BLUE)
doc.add_paragraph().paragraph_format.space_after=Pt(60)
add_table(doc, ['项目概览','内容'], [
    ['产品名称','AI生产排产副驾（Production Planning Copilot）'],
    ['目标客户','以离散制造为主、仍依赖 Excel 与经验排产的中小制造企业'],
    ['核心价值','缩短重排响应、降低交付风险，并建立可复核的排产决策链'],
    ['文档用途','Hackathon 项目路演、客户试点沟通与早期商业化讨论'],
], [1.35,4.95])
doc.add_paragraph('\n版本：V1.0  |  日期：2026年9月', style='Normal').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading('执行摘要', 1)
para(doc, '我们建议将本项目定位为中小制造企业的生产规划副驾：它不取代计划员，也不把大语言模型当成排产权威；而是把订单、物料、设备和人员汇入同一生产快照，由确定性内核生成可行方案，并让计划员在清晰的证据与审批闸门下完成最终决策。')
para(doc, '客户购买的不是“会聊天的 AI”，而是一条从数据导入、约束校验、影响推演、人工审批到可追溯复盘的决策闭环。该闭环尤其适合订单波动、物料短缺、设备停机和人员约束频繁出现，但数字化预算与 IT 人力有限的制造现场。')
add_table(doc, ['商业命题','说明'], [
    ['痛点','手工 Excel 排产耗时、异常重排慢、知识依赖个人且决策难以复盘。'],
    ['解决方案','用确定性排产与约束校验保证可靠性，用 AI 辅助理解脏数据、自然语言场景和解释叙述。'],
    ['采用路径','以单车间试点切入，先证明响应效率、可排产透明度与计划员采纳，再扩展 ERP/MES 集成和多场景模块。'],
    ['商业模式','实施服务 + 订阅许可；按站点、计划员席位和集成复杂度分层报价。'],
], [1.45,4.85])

doc.add_heading('1 业务问题与目标客户', 1)
para(doc, '在大量中小制造车间，计划员每天需要在订单交期、库存、设备能力、工人技能和换型成本之间做取舍。信息通常散落在 Excel、纸质记录或个人经验中。一旦出现插单、缺料或设备停机，原计划需要重新计算，而车间在等待期间缺少统一、可信的答案。')
add_table(doc, ['目标客户画像','典型特征','优先切入理由'], [
    ['离散制造 SME','多品种、小批量；订单与资源约束频繁变化','排产复杂度已超过表格，但尚未具备昂贵 APS 的实施条件'],
    ['生产计划负责人','承担交付、产能协调与异常响应','是直接使用者与采购倡导者，需要可解释、可修改的建议'],
    ['工厂经营者','关注准交、库存占用和人员效率','需要可量化、可审计的业务价值，而非黑箱算法'],
    ['IT/数字化负责人','维护成本敏感，需兼容既有数据流程','可先以 CSV/XLSX 接入，再逐步连接 ERP/MES'],
], [1.4,2.7,2.2])

doc.add_heading('2 产品方案与差异化', 1)
para(doc, '产品以“建议与执行分离”为基本原则。排产、硬约束、评分、基线对比与沙箱计算均由确定性 Python 内核完成；AI 仅服务于表格映射、自然语言场景翻译、解释、风险叙述及受控编排，输出会经过结构化校验。正式计划必须经人工审批后才会生效。')
add_table(doc, ['能力模块','客户得到的结果','当前项目依据'], [
    ['生产快照与排产','订单、物料、机器、人员统一进入排产；输出甘特图与部分可行计划','Dashboard、Schedule 与确定性排产内核'],
    ['约束透明化','每项不可排产作业给出阻塞原因和解锁建议','硬约束校验、PARTIAL 计划'],
    ['What-if 沙箱','在不改动正式计划的前提下模拟停机、缺料等影响','Scenario Sandbox 与计划对比'],
    ['人工审批与审计','审批前重校验输入与约束，保留方案、操作与工具调用轨迹','Approval、Trace 与 Audit Log'],
    ['价值与风险视图','同口径对比 FCFS 基线，展示风险、瓶颈、KPI 与成本口径','Value Ledger、Risks、Insights'],
], [1.35,2.65,2.3])
para(doc, '差异化不在于单点“自动排程”，而在于让计划员敢于使用：系统既说明“为什么这样排”，也诚实说明“为什么不能排”；既允许推演，也将执行权保留给人；既展示预测，也区分实测与预测。')

doc.add_heading('3 客户价值与试点假设', 1)
para(doc, '本项目将价值验证设计成可衡量的试点，而非预先承诺固定收益。基线应使用客户现行排产方式或系统内置 FCFS 同输入对比；试点结束后，由客户与项目团队共同确认口径。')
add_table(doc, ['验证维度','建议指标','试点验证方法'], [
    ['响应效率','异常发生到提交可审阅方案的时间','记录设备停机、缺料、插单等场景的处理时长'],
    ['交付风险','迟交订单数、拖期分钟、零裕度订单','与客户历史计划或 FCFS 基线在同一输入下比较'],
    ['计划透明度','不可排产项是否有明确原因与责任动作','抽样审核阻塞原因、解锁建议和审批记录'],
    ['计划员采纳','建议采纳率、修改原因、人工步骤节省','在审批与偏好规则中记录真实决策'],
    ['治理可信度','输入来源、版本、审批与运行轨迹完整性','检查 Trace、审计日志与导入批次记录'],
], [1.35,2.35,2.6])

doc.add_heading('4 商业模式与定价建议', 1)
para(doc, '建议采用“低门槛试点 + 订阅扩展”的销售路径。下列价格为早期商业化的建议区间，需依据行业、集成范围、交付成本和客户支付意愿验证，不应视作既定报价。')
add_table(doc, ['产品包','适用范围','建议定价方式','包含内容'], [
    ['试点包','单车间、限定数据集、6-10周','一次性实施费，建议人民币 5-15 万','数据盘点、导入映射、场景演练、验收报告'],
    ['基础订阅','单站点、核心计划团队','年费，建议人民币 12-30 万','排产、审批、What-if、风险与基础支持'],
    ['专业订阅','多角色协同、定制指标与集成','年费 + 集成服务，建议人民币 30-80 万起','ERP/MES 接口、价值台账、权限与扩展支持'],
    ['增值服务','复杂工艺或多工厂推广','按项目或人天报价','数据治理、模型接入、流程再设计、培训'],
], [1.25,1.55,1.8,1.7])
para(doc, '收入结构上，实施服务负责覆盖早期数据接入与变更管理成本；订阅收入承接持续使用、产品迭代和客户成功服务。随着标准化导入、模板与接口沉淀，交付时间与边际支持成本应逐步下降。')

doc.add_heading('5 市场进入与销售策略', 1)
para(doc, '早期不宜面向所有制造业泛化销售。应优先聚焦订单波动明显、工序不超过中等复杂度、已有电子表格数据、且管理层能参与试点复盘的离散制造场景，例如机加工、钣金、装配及定制零部件供应商。')
add_table(doc, ['阶段','目标','关键动作','成功信号'], [
    ['阶段一 试点验证','拿到 2-3 个设计合作客户','以真实异常场景跑通数据导入、排产、审批与复盘','客户认可指标口径并愿意提供推荐或续约意向'],
    ['阶段二 可复制交付','形成行业模板与标准实施方法','沉淀数据字典、导入模板、培训包和验收清单','试点周期缩短，交付依赖个性化开发下降'],
    ['阶段三 生态扩展','提升客单价与续费能力','连接 ERP/MES，拓展报价、风险治理与多站点协同','订阅收入占比提升，合作伙伴参与获客与交付'],
], [1.25,1.45,2.45,1.15])
bullets(doc, ['获客路径：行业协会与园区、制造业数字化服务商、ERP/MES 集成商、标杆客户转介绍。', '销售主线：用一个客户真实的停机或缺料场景展示“影响可见、计划不被误改、负责人可审批”。', '采购主张：将项目包装为可控的运营改善试点，而非一次性大规模系统替换。'])

doc.add_heading('6 技术可行性与信任设计', 1)
para(doc, '系统采用 React/Vite 前端、FastAPI 服务层、受控 Agent 与工具闸门、纯 Python 排产内核和 SQLite 持久层的分层架构。部署可从单实例开始，后续保留向 PostgreSQL 等生产级组件迁移的路径。')
add_table(doc, ['风险点','产品控制方式','商业意义'], [
    ['AI 输出不可靠','LLM 不拥有排产与约束判定权；结构化校验与确定性回退','降低客户对黑箱自动化的顾虑'],
    ['错误改动生产计划','沙箱只读推演；正式计划必须走服务端审批与重校验','适合高责任的车间决策流程'],
    ['脏数据与提示注入','导入映射、歧义确认、不可信内容隔离、工具白名单','能从现实表格开始，而非要求完美主数据'],
    ['成本与可用性','token 预算、调用上限、REPLAY 与确定性降级模式','控制 AI 使用成本，模型不可用时核心仍可运行'],
    ['决策无法复盘','Trace、不可篡改审计与价值台账','为客户内部管理与持续采购提供证据'],
], [1.3,3.0,1.95])

doc.add_heading('7 实施路线与资源需求', 1)
add_table(doc, ['周期','重点交付','客户投入','项目团队投入'], [
    ['第1-2周','业务访谈、数据盘点、试点 KPI 与验收口径','计划、工艺、库存和设备数据负责人参与','解决方案负责人、数据实施人员'],
    ['第3-4周','数据映射、种子场景与约束核验','提供样例表格并确认字段与异常处理规则','产品/工程团队配置导入与规则'],
    ['第5-6周','排产、审批、What-if 与风险演练','计划员参与真实或回放场景验证','实施顾问、前后端与算法支持'],
    ['第7-8周','复盘、价值台账、推广决策','确认试点结果、续约与扩展范围','客户成功与商业负责人'],
], [1.05,2.25,1.55,1.4])
para(doc, '建议的最小商业化团队配置为：产品/行业负责人 1 名、全栈工程师 1-2 名、排产或制造领域顾问 1 名、客户成功/实施角色 1 名。早期重点不是扩大销售队伍，而是证明试点可复制，并把每次客户实施沉淀为可复用的模板。')

doc.add_heading('8 风险与应对', 1)
add_table(doc, ['商业风险','可能影响','应对措施'], [
    ['客户数据质量低','影响初始排产结果与信任建立','从 CSV/XLSX 映射与人工确认切入；先限定试点数据范围'],
    ['实施过度定制','拉长交付周期，压低毛利','建立行业模板与变更边界；把集成作为单独服务项'],
    ['价值难量化','采购决策和续约缺乏证据','在启动前冻结 KPI 口径，并使用同输入基线对比'],
    ['现场不愿授权 AI','采纳率低','强调副驾而非替代；保留审批闸门、可见证据和可关闭规则'],
    ['产品能力边界被误解','承诺超出当前交付能力','销售材料明确区分当前页面能力、API 能力与路线图能力'],
], [1.35,2.3,2.6])

doc.add_heading('9 未来12个月里程碑', 1)
add_table(doc, ['时间窗口','产品与商业目标','验收标准'], [
    ['0-3个月','完成产品化试点版本与首批设计合作客户','至少跑通一条真实数据到审批复盘的完整闭环'],
    ['4-6个月','建立标准试点包、行业数据模板与报价方法','试点可在约 6-10 周完成，并形成可复用案例'],
    ['7-9个月','深化 ERP/MES 集成和价值台账；验证续约','至少一个客户从试点进入订阅或扩展阶段'],
    ['10-12个月','形成垂直行业解决方案与伙伴渠道','交付方法、接口与客户成功流程具备规模化基础'],
], [1.2,3.2,1.85])

doc.add_heading('结论与合作请求', 1)
para(doc, 'AI生产排产副驾的商业机会，在于把制造现场最难被数字化的一段工作——跨订单、资源和异常的实时权衡——变成一个可信、可解释、可审批的流程。项目已经具备完整的演示闭环和清晰的技术边界，下一步应通过真实车间试点验证数据适配、价值口径与客户付费意愿。')
para(doc, '我们建议与首批设计合作客户共同启动单车间试点：选择一个包含设备停机、物料短缺或加急插单的高频场景，在既定 KPI 下验证方案响应速度、风险透明度、计划员采纳和治理证据。试点成功后，再按站点、模块和集成范围进入订阅与扩展合作。')

# footer
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('AI生产排产副驾商业计划书  |  Hackathon 项目'); r.font.size = Pt(8); r.font.color.rgb = RGBColor(90,90,90)

doc.core_properties.title = 'AI生产排产副驾商业计划书'
doc.core_properties.subject = 'AI Production Planning Copilot business proposal'
doc.save(OUT)
print(OUT)
